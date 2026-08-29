"""SiCoPR config editor -- a local web app for building and inspecting configs.

    python gui/app.py                 # then open http://127.0.0.1:8765

Deliberately stdlib-only. The engine already needs numpy, scipy and openpyxl;
a GUI that added Flask or FastAPI would make the repo harder to clone and run,
and this needs to serve a handful of JSON endpoints to one local user.

Binds to 127.0.0.1 only. This reads and writes files anywhere the path parameter
points, so it must not be reachable from the network; `_safe_path` additionally
refuses anything outside the repo.

This also SPAWNS PROCESSES (`/api/exec/start`). The command is built as an argv
list from validated repo-relative paths and handed to subprocess without a
shell, and only one run may be live at a time.

Endpoints
    GET  /                     the page
    GET  /api/configs          configs found in the repo
    GET  /api/config?path=     keywords, packages, schema, schematic blocks
    POST /api/save             {src, dst, changes} -> {ok, warnings}
    GET  /api/results          result directories, newest output first
    GET  /api/run?path=        one run: cases, metrics, figures by stage
    GET  /figure?path=         a generated .png
    GET  /api/channel_dirs     directories containing Touchstone files
    GET  /api/channels?dir=    the .s4p files in one, with a guessed role
    POST /api/netlist          a selection -> the exact argv, or the problems
    POST /api/exec/start       run sicopr.py with that argv
    GET  /api/exec/status      streamed output from `since` onward
    POST /api/exec/stop        terminate the current run
"""
import glob
import json
import os
import subprocess
import sys
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
sys.path.insert(0, _ROOT)

from gui import runner, schematic                                   # noqa: E402
from gui.config_io import (read_config, write_config, all_cells,  # noqa: E402
                           split_tx_rx, join_tx_rx, is_tx_rx,
                           engine_keywords, is_setting)

SHEET_NAME = "COM_Settings"
STATIC = os.path.join(_HERE, "static")
PORT = int(os.environ.get('SICOPR_GUI_PORT', '8765'))

CONFIG_GLOBS = [
    'tests/1_IEEE_802p3dj_COM_Spreadsheets/*.xlsx',
    'tests/configs/*.xlsx',
    'config/*.xlsx',
    'configs/*.xlsx',
]
RESULT_GLOBS = ['results*/', 'results/*/']


def _safe_path(p):
    """Absolute path inside the repo, or None.

    The browser is local and trusted, but a path parameter that can walk out of
    the repo turns a convenience server into a file-read primitive for anything
    else running on the machine.
    """
    if not p:
        return None
    full = os.path.abspath(os.path.join(_ROOT, p))
    if os.path.commonpath([full, _ROOT]) != _ROOT:
        return None
    return full


def _jsonable(v):
    """Cell values are floats, ints, strings or None. Anything else is a bug in
    the reader, and stringifying it is more useful than a 500."""
    if v is None or isinstance(v, (bool, int, float, str)):
        if isinstance(v, float) and (v != v or v in (float('inf'), float('-inf'))):
            return str(v)
        return v
    return str(v)


def _cell_json(name, kw, schema, half=None, engine_read=True):
    """One editable field.

    `half` is 'tx' or 'rx' for a `[TX RX]` cell being shown under one end of the
    link; the field then edits that half and the server recombines on save.
    """
    s = schema.get(name, {})
    pair = split_tx_rx(kw.value) if half else None
    out = {
        'half': half,
        'tx_rx': bool(half),
        # the key the client sends back; a half-field is recombined server-side
        'key': ('%s#%s' % (name, half)) if half else name,
        'engine_read': engine_read,
        'half_value': (pair[0] if half == 'tx' else pair[1]) if pair else None,
    }
    out.update({
        'name': name,
        'value': _jsonable(kw.value),
        'units': _jsonable(kw.units),
        'note': _jsonable(kw.info),
        'coord': kw.coord,
        'is_formula': kw.is_formula,
        'formula': kw.formula,
        # A value that is a string is a MATLAB range: a sweep, not a number.
        # The UI must show this, because collapsing a sweep to its first point
        # is a silent 198x reduction of the search space.
        'is_sweep': isinstance(kw.value, str) and any(
            c in str(kw.value) for c in ':['),
        'matlab': s.get('matlab', ''),
        'default': _jsonable(s.get('default')),
        'info': s.get('info', ''),
    })
    # A half-field is a sweep only if its own half is a range, not because the
    # combined literal contains a bracket.
    if half:
        hv = str(out['half_value'] or '')
        out['is_sweep'] = ':' in hv
    return out


def _looks_like_config(path):
    """Cheap test: does this workbook have a parameter sheet the engine reads?

    The repo holds plenty of .xlsx that are not configs -- MATLAB result
    workbooks, comparison tables. Offering those in the config picker and then
    failing to parse them would be worse than not listing them.
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return SHEET_NAME in wb.sheetnames or any(
                s.lower().startswith('keywords') for s in wb.sheetnames)
        finally:
            wb.close()
    except Exception:                                 # noqa: BLE001
        return False


def list_config_dirs():
    """Every directory holding at least one configuration workbook."""
    seen = {}
    for root, dirs, files in os.walk(_ROOT):
        dirs[:] = [d for d in dirs
                   if d not in ('.git', '__pycache__', '.venv', 'node_modules')]
        n = 0
        for f in files:
            if f.lower().endswith('.xlsx') and not f.startswith('~$') \
                    and _looks_like_config(os.path.join(root, f)):
                n += 1
        if n:
            seen[os.path.relpath(root, _ROOT).replace('\\', '/')] = n
    return [{'dir': k, 'count': v} for k, v in sorted(seen.items())]


def list_configs(rel_dir=None):
    """Configs in one directory, or in the conventional locations."""
    paths = []
    if rel_dir:
        full = _safe_path(rel_dir)
        if not full or not os.path.isdir(full):
            raise FileNotFoundError(rel_dir)
        paths = sorted(glob.glob(os.path.join(full, '*.xlsx')))
    else:
        for g in CONFIG_GLOBS:
            paths += sorted(glob.glob(os.path.join(_ROOT, g)))
    out = []
    for p in paths:
        if os.path.basename(p).startswith('~$') or not _looks_like_config(p):
            continue
        out.append({'path': os.path.relpath(p, _ROOT).replace('\\', '/'),
                    'name': os.path.basename(p),
                    'size': os.path.getsize(p)})
    return out


def _scalarish(v):
    """True for a plain number, False for a vector, matrix or sweep range."""
    try:
        float(str(v).strip())
        return True
    except (TypeError, ValueError):
        return False


def _shape(v):
    """'num', 'bracketed' or 'bare' -- how the engine will read this literal.

    Bracketing is not cosmetic. `C_d` declares its default as
    `4e-05 4e-05 9e-05 ...` while the config holds
    `[0.4e-4  0.9e-4  1.1e-4 ; ...]`; both are "not a number", but the bare
    form reaches the engine as a plain string and the first arithmetic on it
    raises `can't multiply sequence by non-int of type 'float'`. Measured: of
    47 shape-compatible defaults, that one alone made the whole config
    unreadable.
    """
    if _scalarish(v):
        return 'num'
    t = str(v).strip()
    return 'bracketed' if t.startswith('[') and t.endswith(']') else 'bare'


def new_config(template, dst_rel, reset_defaults=True):
    """Create a config from a template, optionally reset to schema defaults.

    There is no way to build a runnable config from nothing: the engine needs
    the `.START` package blocks, and the workbook's own `keywords_*` sheet is
    what documents it. So "new" means "a copy of a known-good workbook with the
    main-region settings put back to the defaults the workbook itself
    declares", which is reproducible and cannot produce a file the engine
    refuses to read.

    Package blocks and formula cells are left alone -- a default of 0 for a
    package geometry would be a broken config, not a blank one.
    """
    src = _safe_path(template)
    dst = _safe_path(dst_rel)
    if not src or not os.path.isfile(src):
        return {'error': 'template not found: %r' % template}
    if not dst or not dst.lower().endswith('.xlsx'):
        return {'error': 'destination must be a .xlsx path inside the repo'}
    if os.path.exists(dst):
        return {'error': 'refusing to overwrite %s'
                         % os.path.relpath(dst, _ROOT), 'exists': True}

    changes = {}
    skipped = []
    if reset_defaults:
        doc = read_config(src)
        for name, kw in doc.keywords.items():
            if not is_setting(name) or kw.is_formula:
                continue
            d = doc.schema.get(name, {}).get('default')
            # Not every "default" in the keywords sheet is a usable value.
            # `PKG_NAME` declares `-`, a placeholder meaning "name a package
            # block"; writing it produces a workbook the engine refuses with
            # `Package Block "-" not found`. A placeholder is not a default.
            if d is None or str(d).strip() in ('', '-', '--', 'n/a', 'N/A',
                                               'none', 'None', 'TBD'):
                skipped.append(name)
                continue
            if str(d).strip() == str(kw.value).strip():
                continue
            # The keywords sheet documents a default PER KEYWORD; the set of
            # them is not a validated configuration. Several are scalars where
            # the config holds a vector or a sweep -- `C_b` defaults to `0`
            # against a `[TX RX]` pair, and `g_DC` to `-52` (itself a stray
            # formula in that sheet) against a `[-20:1:0]` range. Writing those
            # produces a workbook the engine cannot read at all.
            #
            # So a default is only applied when it is the same SHAPE as what it
            # replaces: scalar for scalar, literal for literal.
            if _shape(d) != _shape(kw.value):
                skipped.append(name)
                continue
            changes[name] = str(d).strip()

    os.makedirs(os.path.dirname(dst), exist_ok=True)
    warnings = write_config(src, dst, changes)

    # Hand back nothing the engine cannot read. Resetting ~55 settings at once
    # is exactly where a bad default turns into a config that only fails much
    # later, in the middle of a run.
    ok, why = _engine_accepts(dst)
    if not ok:
        os.remove(dst)
        return {'error': 'the config built from those defaults is not one the '
                         'engine can read, so it was not kept: %s' % why}

    return {'ok': True, 'dst': os.path.relpath(dst, _ROOT).replace('\\', '/'),
            'reset': len(changes), 'no_default': len(skipped),
            'warnings': warnings}


_VALIDATE_SRC = r'''
import sys
from types import SimpleNamespace
sys.path.insert(0, %r)
import sicopr
lines = open(%r, encoding='utf-8').read().splitlines()
a = next(i for i, l in enumerate(lines) if l.strip() == 'OP = SimpleNamespace()')
b = next(i for i, l in enumerate(lines) if 'param, OP = read_ParamConfigFile' in l)
ns = {'SimpleNamespace': SimpleNamespace}
exec(chr(10).join(l[4:] for l in lines[a:b]), ns)
sicopr.read_ParamConfigFile(sys.argv[1], ns['OP'])
print('OK')
'''


def _engine_accepts(path):
    """(True, '') if `read_ParamConfigFile` parses this workbook.

    Run in a subprocess for two reasons. It keeps the GUI independent of the
    engine -- no heavy import at startup, and the editor still runs while the
    engine is mid-edit. And it avoids a file lock: the engine's Excel loader
    opens the workbook with `read_only=True` and never closes it
    (`__load_excel`), so on Windows a config validated in-process cannot then
    be deleted or replaced.
    """
    engine = os.path.join(_ROOT, 'sicopr.py')
    if not os.path.isfile(engine):
        return True, ''            # nothing to validate against
    src = _VALIDATE_SRC % (_ROOT, engine)
    try:
        p = subprocess.run([sys.executable, '-c', src, path],
                           cwd=_ROOT, capture_output=True, text=True,
                           timeout=180)
    except subprocess.TimeoutExpired:
        return False, 'the engine took too long to parse it'
    if p.returncode == 0:
        return True, ''
    tail = [l for l in (p.stderr or p.stdout).strip().split('\n') if l.strip()]
    return False, tail[-1] if tail else 'exit %s' % p.returncode


def config_payload(rel):
    full = _safe_path(rel)
    if not full or not os.path.isfile(full):
        raise FileNotFoundError(rel)
    doc = read_config(full)
    schema = doc.schema

    # Most cells the engine does not read are the units column and the
    # `[min:step:max]` templates -- and those are ALREADY carried on the
    # setting they annotate, as its `units` and `note`. Collecting them into a
    # block of their own just showed the same text twice, once out of context.
    #
    # So an annotation is dropped only when its text provably appears on a real
    # setting; the rest (reference values like T_t / T_ft / T_nt, which neither
    # the port nor MATLAB reads) stay, in the block they belong to, flagged as
    # not read by the engine. Nothing is silently lost either way.
    settings = [n for n in doc.keywords if is_setting(n)]
    annotations = [n for n in doc.keywords if not is_setting(n)]

    # An annotation is redundant if its text is already on screen as some real
    # setting's value, units or note. Values matter here as much as units: a
    # value cell holding a string (`[0.13 0.15 ; 0.13 0.15]`, `[-20:1:0]`) has
    # the units cell to ITS right, so the reader picks it up as a label too and
    # it came back as a phantom setting whose "value" was `nH`.
    shown_text = set()
    for n in settings:
        kw = doc.keywords[n]
        for x in (kw.value, kw.units, kw.info):
            if x is not None:
                shown_text.add(str(x).strip())
    HEADERS = {'parameter', 'setting', 'units', 'information', 'info'}
    redundant = {n for n in annotations
                 if n.strip() in shown_text or n.strip().lower() in HEADERS}
    kept = [n for n in annotations if n not in redundant]

    blocks = schematic.assign(settings + kept)
    for b in blocks:
        cells = []
        for n in b['keywords']:
            kw = doc.keywords[n]
            paired = is_tx_rx(kw)
            if paired and b['id'] in schematic.MIRROR:
                cells.append(_cell_json(n, kw, schema, half='tx'))
            else:
                cells.append(_cell_json(n, kw, schema,
                                        engine_read=is_setting(n)))
        b['cells'] = cells

    # mirror the RX half of every [TX RX] cell into the paired Rx-side block
    by_id = {b['id']: b for b in blocks}
    for tx_id, rx_id in schematic.MIRROR.items():
        src, dst = by_id.get(tx_id), by_id.get(rx_id)
        if not src or not dst:
            continue
        for n in src['keywords']:
            kw = doc.keywords[n]
            if is_tx_rx(kw):
                dst['cells'].append(_cell_json(n, kw, schema, half='rx'))
    for b in blocks:
        del b['keywords']

    packages = {}
    for pkg, cells in doc.packages.items():
        packages[pkg] = [_cell_json(n, c, schema) for n, c in sorted(cells.items())]

    selected = doc.keywords.get('PKG_NAME')
    return {
        'path': rel,
        'name': os.path.basename(full),
        'blocks': blocks,
        'packages': packages,
        'pkg_selected': _jsonable(selected.value) if selected else '',
        'counts': {'main': len(doc.keywords),
                   'settings': len(settings),
                   'annotations': len(annotations),
                   'annotations_folded': len(redundant),
                   'packages': sum(len(v) for v in doc.packages.values()),
                   'addressable': len(all_cells(doc)),
                   'engine_total': len(engine_keywords())},
    }


def list_results():
    out = []
    for g in RESULT_GLOBS:
        for p in glob.glob(os.path.join(_ROOT, g)):
            if not os.path.isdir(p):
                continue
            rel = os.path.relpath(p, _ROOT).replace('\\', '/')
            csvs = glob.glob(os.path.join(p, '**', '*.csv'), recursive=True)
            pngs = glob.glob(os.path.join(p, '**', '*.png'), recursive=True)
            cases = [d for d in os.listdir(p)
                     if d.startswith('case_') and os.path.isdir(os.path.join(p, d))]
            # `results*/` also matches the `results/` container itself, which
            # would otherwise appear as a run aggregating every run inside it.
            # A run has either case directories or output files of its own.
            direct = glob.glob(os.path.join(p, '*.csv')) + \
                glob.glob(os.path.join(p, '*.png'))
            if not cases and not direct:
                continue
            # Sort by the newest FILE inside, not the directory's own mtime.
            # A directory's mtime does not change when a file in one of its
            # sub-directories is rewritten, and per-case output lands in
            # case_NN/. Configs commonly set a date-templated RESULT_DIR, so a
            # re-run overwrites an existing directory -- and using the directory
            # mtime would leave the run that just finished sorted below runs
            # that are days older, which is precisely wrong for "refresh to see
            # the latest results".
            newest = os.path.getmtime(p)
            for f in csvs + pngs:
                try:
                    newest = max(newest, os.path.getmtime(f))
                except OSError:
                    pass
            out.append({'path': rel, 'mtime': newest,
                        'csv': len(csvs), 'png': len(pngs), 'cases': len(cases)})
    out.sort(key=lambda r: -r['mtime'])
    return out[:50]


def _expand_half_keys(src, changes):
    """Turn `C_d#tx` / `C_d#rx` edits back into one write of the whole cell.

    The Tx and Rx blocks each edit their own end of a `[TX RX]` cell, but the
    engine reads a single cell, so the two halves are recombined here -- in the
    shape the original literal had. `write_config` never learns about halves,
    which keeps the tested writer exactly as it was.

    Editing one half takes the other from the file, so a half-edit cannot
    silently blank the end the user did not touch.
    """
    if not any('#' in k for k in changes):
        return changes

    doc = read_config(src)
    out = {k: v for k, v in changes.items() if '#' not in k}
    halves = {}
    for k, v in changes.items():
        if '#' not in k:
            continue
        name, _, side = k.rpartition('#')
        if side not in ('tx', 'rx'):
            raise KeyError('unknown half %r in %r' % (side, k))
        if name not in doc.keywords:
            raise KeyError('not addressable keyword: %r' % name)
        halves.setdefault(name, {})[side] = v

    for name, sides in halves.items():
        original = doc.keywords[name].value
        pair = split_tx_rx(original)
        if pair is None:
            raise KeyError('%r does not hold a [TX RX] pair (%r); it cannot be '
                           'edited one half at a time' % (name, original))
        tx = sides.get('tx', pair[0])
        rx = sides.get('rx', pair[1])
        out[name] = join_tx_rx(tx, rx, original)
    return out


def _guess_role(name):
    """thru / fext / next, from the filename.

    A convenience for pre-selecting, never a decision: the netlist shows what
    was chosen and the user can move any file to any role. Guessing silently
    would be how a NEXT aggressor ends up analysed as the victim channel.
    """
    n = name.lower()
    if 'fext' in n:
        return 'fext'
    if 'next' in n:
        return 'next'
    if 'thru' in n or 'through' in n:
        return 'thru'
    return 'unknown'


def list_channel_dirs():
    """Directories holding Touchstone files, with a count each."""
    seen = {}
    for root, dirs, files in os.walk(_ROOT):
        dirs[:] = [d for d in dirs
                   if d not in ('.git', '__pycache__', '.venv', 'node_modules')]
        s4p = [f for f in files if f.lower().endswith('.s4p')]
        if s4p:
            rel = os.path.relpath(root, _ROOT).replace('\\', '/')
            seen[rel] = len(s4p)
    return [{'dir': k, 'count': v} for k, v in sorted(seen.items())]


def list_channels(rel_dir):
    full = _safe_path(rel_dir)
    if not full or not os.path.isdir(full):
        raise FileNotFoundError(rel_dir)
    out = []
    for f in sorted(os.listdir(full)):
        if not f.lower().endswith('.s4p'):
            continue
        p = os.path.join(full, f)
        out.append({
            'path': os.path.relpath(p, _ROOT).replace('\\', '/'),
            'name': f,
            'role': _guess_role(f),
            'size': os.path.getsize(p),
        })
    return out


def build_netlist(body):
    """Resolve a GUI selection into the exact argv `sicopr.py` will be given.

    Returns (argv, problems). Problems are reported rather than raised so the
    page can show the whole netlist with every issue marked, instead of the
    first one only.
    """
    problems = []

    def resolve(rel, kind, ext):
        full = _safe_path(rel)
        if not full:
            problems.append('%s: %r is outside the repository' % (kind, rel))
            return None
        if not full.lower().endswith(ext):
            problems.append('%s: %r is not a %s file' % (kind, rel, ext))
            return None
        if not os.path.isfile(full):
            problems.append('%s: %r does not exist' % (kind, rel))
            return None
        return full

    config = body.get('config') or ''
    thru = body.get('thru') or ''
    if not config:
        problems.append('no configuration workbook selected')
    if not thru:
        problems.append('no THRU (victim) channel selected')

    cfg_full = resolve(config, 'config', '.xlsx') if config else None
    thru_full = resolve(thru, 'thru', '.s4p') if thru else None

    fext, nxt = [], []
    for rel in body.get('fext') or []:
        f = resolve(rel, 'fext', '.s4p')
        if f:
            fext.append(f)
    for rel in body.get('next') or []:
        f = resolve(rel, 'next', '.s4p')
        if f:
            nxt.append(f)

    if problems:
        return None, problems

    argv = [sys.executable, os.path.join(_ROOT, 'sicopr.py'), cfg_full, thru_full]
    if fext:
        argv += ['--fext'] + fext
    if nxt:
        argv += ['--next'] + nxt
    if body.get('export_mat'):
        argv.append('--export-mat')
    if body.get('eye_under_mlse'):
        argv.append('--eye-under-mlse')
    mv = body.get('matlab_version')
    if mv in ('4p15p0', '4p16p0'):
        argv += ['--matlab-version', mv]
    return argv, []


def _display_argv(argv):
    """The command as a person would type it — repo-relative, python first."""
    out = ['python']
    for a in argv[1:]:
        if os.path.isabs(a) and a.startswith(_ROOT):
            a = os.path.relpath(a, _ROOT).replace('\\', '/')
        out.append('"%s"' % a if ' ' in a else a)
    return ' '.join(out)


# The values worth putting in front of someone opening a run, in the order a
# reader wants them: the answer first, then what produced it.
HEADLINE = [
    ('COM_dB', 'COM', 'dB'),
    ('VEO_mV', 'VEO', 'mV'),
    ('VEC_dB', 'VEC', 'dB'),
    ('ERL', 'ERL', 'dB'),
    ('FOM', 'FOM', 'dB'),
    ('itick', 'itick', ''),
    ('CTLE_DC_gain_dB', 'CTLE g_DC', 'dB'),
    ('g_DC_HP', 'g_DC_HP', 'dB'),
    ('TXLE_taps', 'Tx FFE', ''),
    ('DFE_taps', 'DFE', ''),
    ('available_signal_after_eq_mV', 'A_s', 'mV'),
    ('sigma_N', 'sigma_N', ''),
]

STAGE_NAMES = {
    '1': 'Channel / FD', '2': 'TDR / ERL', '3': 'Pulse (TD)',
    '4': 'Sampling', '5': 'Equalization', '6': 'Noise', '7': 'COM',
}


def _read_case_csv(path):
    """First data row of a per-case results.csv, as {column: text}."""
    import csv
    try:
        with open(path, newline='', encoding='utf-8-sig') as f:
            rows = list(csv.reader(f))
    except OSError:
        return {}
    if len(rows) < 2:
        return {}
    return dict(zip(rows[0], rows[1]))


def run_payload(rel):
    """One results directory: its cases, their headline numbers, and the
    figures each case produced, grouped by pipeline stage."""
    full = _safe_path(rel)
    if not full or not os.path.isdir(full):
        raise FileNotFoundError(rel)

    cases = []
    names = sorted(d for d in os.listdir(full)
                   if os.path.isdir(os.path.join(full, d)))
    # a run with no case_NN subdirectories may itself be one case
    if not names:
        names = ['']

    for name in names:
        cdir = os.path.join(full, name) if name else full
        row = _read_case_csv(os.path.join(cdir, 'results.csv'))
        metrics = []
        for col, label, unit in HEADLINE:
            if col in row and str(row[col]).strip() != '':
                metrics.append({'label': label, 'unit': unit,
                                'value': str(row[col]).strip()})

        stages = {}
        for png in sorted(glob.glob(os.path.join(cdir, '*.png'))):
            base = os.path.basename(png)
            sid = base[1] if base.startswith('s') and base[1:2].isdigit() else '?'
            stages.setdefault(sid, []).append(
                os.path.relpath(png, _ROOT).replace('\\', '/'))

        cases.append({
            'name': name or os.path.basename(full),
            'metrics': metrics,
            'has_csv': bool(row),
            'columns': len(row),
            # every column of results.csv, so the full output is reachable
            # without leaving the page; the client keeps it collapsed
            'all': [{'name': k, 'value': str(v)} for k, v in row.items()],
            'stages': [{'id': s, 'label': STAGE_NAMES.get(s, 'Other'),
                        'figures': stages[s]}
                       for s in sorted(stages)],
        })

    return {'path': rel, 'name': os.path.basename(full), 'cases': cases}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):        # quieter console
        pass

    def _send(self, code, body, ctype='application/json'):
        if not isinstance(body, bytes):
            body = body.encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', ctype)
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _json(self, obj, code=200):
        self._send(code, json.dumps(obj))

    def do_GET(self):
        u = urlparse(self.path)
        q = parse_qs(u.query)
        try:
            if u.path in ('/', '/index.html'):
                with open(os.path.join(STATIC, 'index.html'), 'rb') as f:
                    return self._send(200, f.read(), 'text/html; charset=utf-8')
            if u.path.startswith('/static/'):
                name = os.path.basename(u.path)
                full = os.path.join(STATIC, name)
                if not os.path.isfile(full):
                    return self._send(404, b'not found', 'text/plain')
                ctype = ('text/css' if name.endswith('.css')
                         else 'application/javascript' if name.endswith('.js')
                         else 'application/octet-stream')
                with open(full, 'rb') as f:
                    return self._send(200, f.read(), ctype + '; charset=utf-8')
            if u.path == '/api/configs':
                return self._json(list_configs(q.get('dir', [None])[0]))
            if u.path == '/api/config_dirs':
                return self._json(list_config_dirs())
            if u.path == '/api/config':
                return self._json(config_payload(q.get('path', [''])[0]))
            if u.path == '/api/results':
                return self._json(list_results())
            if u.path == '/api/run':
                return self._json(run_payload(q.get('path', [''])[0]))
            if u.path == '/api/channel_dirs':
                return self._json(list_channel_dirs())
            if u.path == '/api/channels':
                return self._json(list_channels(q.get('dir', [''])[0]))
            if u.path == '/api/exec/status':
                r = runner.current()
                if r is None:
                    return self._json({'idle': True, 'lines': [], 'next': 0})
                since = int(q.get('since', ['0'])[0] or 0)
                st = r.status(since)
                st['idle'] = False
                st['command'] = _display_argv(r.argv)
                return self._json(st)
            if u.path == '/figure':
                return self._figure(q.get('path', [''])[0])
            return self._send(404, b'not found', 'text/plain')
        except FileNotFoundError as e:
            return self._json({'error': 'not found: %s' % e}, 404)
        except Exception as e:                            # noqa: BLE001
            import traceback
            traceback.print_exc()
            return self._json({'error': '%s: %s' % (type(e).__name__, e)}, 500)

    def _start(self, body):
        argv, problems = build_netlist(body)
        if problems:
            return {'error': 'the inputs are not runnable', 'problems': problems}
        try:
            r = runner.start(argv, _ROOT, label=os.path.basename(body['config']))
        except RuntimeError as e:
            return {'error': str(e)}
        return {'ok': True, 'command': _display_argv(r.argv)}

    def _figure(self, rel):
        """Serve a generated figure. Restricted to .png inside the repo: this
        endpoint exists to show run output, not to read arbitrary files."""
        full = _safe_path(rel)
        if not full or not full.lower().endswith('.png') or not os.path.isfile(full):
            return self._send(404, b'not found', 'text/plain')
        with open(full, 'rb') as f:
            return self._send(200, f.read(), 'image/png')

    def do_POST(self):
        u = urlparse(self.path)
        try:
            n = int(self.headers.get('Content-Length', 0))
            body = json.loads(self.rfile.read(n) or b'{}')
            if u.path == '/api/save':
                return self._json(self._save(body))
            if u.path == '/api/new_config':
                return self._json(new_config(
                    body.get('template', ''), body.get('dst', ''),
                    bool(body.get('reset_defaults', True))))
            if u.path == '/api/netlist':
                argv, problems = build_netlist(body)
                return self._json({
                    'ok': not problems, 'problems': problems,
                    'command': _display_argv(argv) if argv else '',
                    'argv': argv or []})
            if u.path == '/api/exec/start':
                return self._json(self._start(body))
            if u.path == '/api/exec/stop':
                return self._json({'stopped': runner.stop()})
            return self._send(404, b'not found', 'text/plain')
        except Exception as e:                            # noqa: BLE001
            import traceback
            traceback.print_exc()
            return self._json({'error': '%s: %s' % (type(e).__name__, e)}, 500)

    def _save(self, body):
        src = _safe_path(body.get('src'))
        dst = _safe_path(body.get('dst'))
        if not src or not os.path.isfile(src):
            return {'error': 'source config not found: %r' % body.get('src')}
        if not dst:
            return {'error': 'destination is outside the repo: %r' % body.get('dst')}
        if not dst.lower().endswith('.xlsx'):
            return {'error': 'destination must be a .xlsx file'}
        if os.path.exists(dst) and not body.get('overwrite'):
            return {'error': 'refusing to overwrite %s; pass overwrite to force'
                             % os.path.relpath(dst, _ROOT), 'exists': True}
        changes = body.get('changes') or {}
        if not isinstance(changes, dict):
            return {'error': 'changes must be an object'}
        try:
            changes = _expand_half_keys(src, changes)
        except KeyError as e:
            return {'error': str(e).strip('"')}
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        try:
            warnings = write_config(src, dst, changes)
        except KeyError as e:
            # a keyword the workbook does not define is the caller's mistake,
            # not a server fault; report it plainly rather than as a 500
            return {'error': str(e).strip('"')}
        return {'ok': True,
                'dst': os.path.relpath(dst, _ROOT).replace('\\', '/'),
                'changed': len(changes),
                'warnings': warnings}


def main():
    dup = schematic.duplicate_names()
    if dup:
        print('schematic.BLOCKS lists these keywords more than once: %s' % dup)
        return 1
    srv = ThreadingHTTPServer(('127.0.0.1', PORT), Handler)
    url = 'http://127.0.0.1:%d' % PORT
    print('SiCoPR config editor on %s   (Ctrl-C to stop)' % url)
    print('%d config(s) visible' % len(list_configs()))
    if '--no-browser' not in sys.argv:
        try:
            webbrowser.open(url)
        except Exception:                                 # noqa: BLE001
            pass
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print('\nstopped')
    return 0


if __name__ == '__main__':
    sys.exit(main())
