"""Find config keywords whose sweep collapsed to a single point.

The Tx FFE defect (docs/TXFFE_SWEEP_ROOT_CAUSE.md) had a specific shape:

    A            B      C                  D
    c(-1)        0      [ -0.34:.02:0]     [min:step:max]

Both engines read the cell immediately RIGHT of the label (column B). The sweep
range sat one further over, as a template. Nothing errored; the search dimension
simply collapsed to one point and the run looked healthy.

Two checks here:

  ADJACENT-TEMPLATE   a keyword whose value cell is a scalar while a MATLAB
                      range/vector literal sits in the next cell along. That is
                      the exact fingerprint of the Tx FFE defect.

  DEGENERATE-SWEEP    a search dimension that ends up with exactly one point.
                      Legitimate for a pinned run, but it should be a deliberate
                      choice, so it is reported.

    python tools/audit_search_space.py [config.xlsx ...]
"""
import glob
import os
import re
import sys

import numpy as np
import openpyxl

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

# A MATLAB range or vector literal: [a:b:c], [a b c], a:b:c
_RANGE = re.compile(r'^\s*\[?\s*-?[\d.]+\s*:\s*-?[\d.]+\s*(:\s*-?[\d.]+\s*)?\]?\s*$')
_VECTOR = re.compile(r'^\s*\[[^\]]*[\s,][^\]]*\]\s*$')


def looks_like_sweep(v):
    if not isinstance(v, str):
        return False
    s = v.strip()
    if not s or s.lower() in ('[min:step:max]', 'min:step:max'):
        return False          # the literal placeholder text, not a real range
    return bool(_RANGE.match(s) or _VECTOR.match(s))


def scan_sheet(path):
    """Yield (sheet, label, value_cell, neighbour_cell) for the adjacent-template check."""
    wb = openpyxl.load_workbook(path, data_only=True)
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str) or not c.value.strip():
                    continue
                val = ws.cell(c.row, c.column + 1).value
                nxt = ws.cell(c.row, c.column + 2).value
                if val is None:
                    continue
                # value is a plain number, and the next cell along is a real range
                if isinstance(val, (int, float)) and looks_like_sweep(nxt):
                    yield name, c.value.strip(), val, str(nxt).strip()


def search_dims(path):
    """The five EQ search dimensions the optimiser actually loops over."""
    import io as _io
    from types import SimpleNamespace
    import com
    src = _io.open(os.path.join(_ROOT, 'com.py'), encoding='utf-8').read().split('\n')
    a = next(i for i, l in enumerate(src) if l.strip() == 'OP = SimpleNamespace()')
    b = next(i for i, l in enumerate(src) if 'param, OP = read_ParamConfigFile' in l)
    ns = {'SimpleNamespace': SimpleNamespace}
    exec('\n'.join(l[4:] for l in src[a:b]), ns)
    param, OP = com.read_ParamConfigFile(path, ns['OP'])

    def n(v):
        return int(np.atleast_1d(np.asarray(v)).size)

    tsar = np.atleast_1d(np.asarray(param.ts_sample_adj_range)).ravel()
    txffe, _cur, _sw, _fv, _cv = com.OptFom_Build_TXFFE(param)
    return {
        'cursor_gain (Gffe)': n(getattr(param, 'cursor_gain', 0)),
        'ctle_gdc_values': n(param.ctle_gdc_values),
        'g_DC_HP_values': n(getattr(param, 'g_DC_HP_values', 0)),
        'Tx FFE grid': int(np.atleast_2d(txffe).shape[0]),
        'itick range': int(tsar[1] - tsar[0] + 1) if tsar.size >= 2 else 1,
    }


def main():
    paths = sys.argv[1:] or sorted(glob.glob(os.path.join(
        _ROOT, 'tests', '1_IEEE_802p3dj_COM_Spreadsheets', '*.xlsx')))
    bad = 0
    for p in paths:
        print('=== %s' % os.path.basename(p))
        hits = list(scan_sheet(p))
        if hits:
            print('  ADJACENT-TEMPLATE  (scalar value with a range in the next cell):')
            for sheet, label, val, nxt in hits:
                print('     [%s] %-24s value=%-8s next=%s' % (sheet, label, val, nxt))
                bad += 1
        else:
            print('  ADJACENT-TEMPLATE  none')
        try:
            dims = search_dims(p)
        except Exception as exc:
            print('  search dims: could not evaluate (%r)' % (exc,))
            continue
        degen = [k for k, v in dims.items() if v == 1]
        print('  search dimensions: ' + '  '.join('%s=%d' % (k, v) for k, v in dims.items()))
        if degen:
            print('  DEGENERATE-SWEEP   %s' % ', '.join(degen))
    print('\n%d adjacent-template finding(s)' % bad)
    return 0


if __name__ == '__main__':
    sys.exit(main())
